"""Tests for applying a split plan to the docx XML.

The segmenter decides *where* to cut; this layer performs the cut on real
``<w:p>`` elements. The properties worth protecting are that the result
lines up with the model (the applier matches paragraphs positionally, so a
drift means styling the wrong paragraph), that run formatting survives, and
that paragraphs carrying objects are never touched.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt

from papergroom.models import Document, Paragraph
from papergroom.parser import parse_document
from papergroom.structure import structurize
from papergroom.structure.apply import apply_splits
from papergroom.writer import write_document

MERGED = "1 研究背景 当前高校课程论文格式规范问题很多。"


def _build(path: Path, text: str = MERGED, *, bold: bool = True) -> Path:
    doc = DocxDocument()
    run = doc.add_paragraph().add_run(text)
    run.bold = bold
    run.font.size = Pt(16)
    doc.save(str(path))
    return path


def test_splits_reproduce_the_model(tmp_path: Path):
    """The XML paragraph list must match the model's, position for position."""
    path = _build(tmp_path / "merged.docx")
    model = structurize(parse_document(path))

    docx_doc = DocxDocument(str(path))
    apply_splits(docx_doc, model.splits)

    assert [p.text for p in docx_doc.paragraphs] == [p.text for p in model.paragraphs]
    assert [p.text for p in docx_doc.paragraphs] == [
        "1 研究背景",
        "当前高校课程论文格式规范问题很多。",
    ]


def test_run_formatting_survives_a_split_inside_a_run(tmp_path: Path):
    """The cut falls mid-run; both halves must keep the original rPr."""
    path = _build(tmp_path / "merged.docx")
    model = structurize(parse_document(path))

    docx_doc = DocxDocument(str(path))
    apply_splits(docx_doc, model.splits)

    heading, body = docx_doc.paragraphs
    assert heading.runs[0].bold is True
    assert heading.runs[0].font.size == Pt(16)
    # The body's run is a clone of the one that was cut in two.
    assert body.runs[0].bold is True
    assert body.runs[0].font.size == Pt(16)


def test_separator_is_dropped_from_the_output(tmp_path: Path):
    """The colon in "摘要：……" is neither kept nor duplicated."""
    path = _build(tmp_path / "abstract.docx", text="摘要：随着信息化教育快速发展，传统人工排版效率低下。")
    model = structurize(parse_document(path))

    docx_doc = DocxDocument(str(path))
    apply_splits(docx_doc, model.splits)

    assert [p.text for p in docx_doc.paragraphs] == [
        "摘要",
        "随着信息化教育快速发展，传统人工排版效率低下。",
    ]


def test_no_split_leaves_the_document_untouched(tmp_path: Path):
    path = _build(tmp_path / "plain.docx", text="本文研究排版问题，并提出一种自动化方案。")
    model = structurize(parse_document(path))
    assert model.splits == {}

    docx_doc = DocxDocument(str(path))
    apply_splits(docx_doc, model.splits)

    assert [p.text for p in docx_doc.paragraphs] == ["本文研究排版问题，并提出一种自动化方案。"]


def test_paragraph_with_a_drawing_is_never_split(tmp_path: Path):
    """A drawing makes character offsets ambiguous; skip rather than risk it."""
    path = tmp_path / "image.docx"
    doc = DocxDocument()
    para = doc.add_paragraph()
    para.add_run(MERGED)
    para._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:drawing/></w:r>'))
    doc.save(str(path))

    model = structurize(parse_document(path))

    # Refused at the model level...
    assert model.splits == {}
    assert len(model.paragraphs) == 1

    # ...and enforced again at the XML level.
    docx_doc = DocxDocument(str(path))
    apply_splits(docx_doc, {0: [(0, 6), (7, len(MERGED))]})
    assert len(docx_doc.paragraphs) == 1


def test_writer_refuses_when_model_and_file_disagree(tmp_path: Path):
    """A drift between model and file must fail loudly, not mis-style."""
    path = _build(tmp_path / "merged.docx")

    # Claim the paragraph splits, but hand the writer a model that doesn't.
    bogus = Document(
        source_path=str(path),
        paragraphs=[Paragraph(index=0, text=MERGED, source_index=0)],
        splits={0: [(0, 6), (7, len(MERGED))]},
    )

    with pytest.raises(RuntimeError, match="out of sync"):
        write_document(bogus, None, tmp_path / "out.docx")


def test_writer_end_to_end_round_trip(tmp_path: Path):
    """parse -> structurize -> write produces a readable, correctly split file."""
    path = _build(tmp_path / "merged.docx")
    model = structurize(parse_document(path))

    out = write_document(model, None, tmp_path / "out.docx")

    written = DocxDocument(str(out))
    assert [p.text for p in written.paragraphs] == [
        "1 研究背景",
        "当前高校课程论文格式规范问题很多。",
    ]
