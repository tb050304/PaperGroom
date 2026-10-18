"""Apply resolved styles to a python-docx document.

In the new architecture, the AI labeler decides per-paragraph styles
directly (stored on ``Paragraph.resolved_style``). The applier simply
walks the document and applies each paragraph's resolved style.

Key principle: ONLY properties explicitly set by the AI are changed.
All other properties keep the original document formatting.
"""

from __future__ import annotations

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

from ..models import Document, Paragraph, ParagraphStyle, Role

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

# Aliases for Word's built-in styles, used when ``--word-styles`` is on.
#
# Neither the name nor the id is portable on its own. The *name* is localized
# ("Heading 1" / "标题 1"), and the *id* varies by who produced the document:
# a file from a converter can call Normal "a" and Heading 1 "1", while Word's
# own template uses "Normal" and "Heading1". So match on a space- and
# case-insensitive form of the name, accepting both languages, then use
# whatever id that document actually assigned.
_WORD_STYLE_ALIASES: dict[Role, tuple[str, ...]] = {
    Role.DOCUMENT_TITLE: ("title", "标题"),
    Role.TITLE_L1: ("heading 1", "标题 1", "标题1"),
    Role.TITLE_L2: ("heading 2", "标题 2", "标题2"),
    Role.TITLE_L3: ("heading 3", "标题 3", "标题3"),
}


def apply_ruleset(
    doc: Document,
    ruleset,
    docx_doc: DocxDocument,
    *,
    word_styles: bool = False,
) -> None:
    """Apply resolved styles to ``docx_doc`` based on the labeled ``doc``.

    ``ruleset`` is kept for backward compatibility but is NOT used here —
    styles come from ``para.resolved_style``.

    Mutates ``docx_doc`` in place. Caller saves it afterwards.
    Paragraphs in ``doc.paragraphs`` are matched to ``docx_doc.paragraphs``
    by index.

    With ``word_styles=True``, headings also get Word's built-in Heading
    styles, so the navigation pane and an automatic table of contents work.
    Direct formatting is applied afterwards, so the ruleset still decides how
    the heading looks.
    """
    docx_paragraphs = docx_doc.paragraphs
    for para_model in doc.paragraphs:
        if para_model.index >= len(docx_paragraphs):
            break
        para_docx = docx_paragraphs[para_model.index]

        # Apply the semantic style first; direct formatting goes on top.
        styled_as_word = word_styles and _apply_word_style(
            para_docx, para_model.role, docx_doc
        )

        style = para_model.resolved_style

        # No resolved style → change nothing (preserve original formatting).
        if style is None or style.is_empty():
            continue

        # A Word heading style carries a theme colour (Heading 1 is blue in
        # the default template). If the ruleset says nothing about colour,
        # neutralise it rather than shipping blue headings.
        force_black = styled_as_word and style.color_hex is None

        # Protected: formulas only get outer paragraph style, no run edits.
        if para_model.has_formula:
            _apply_paragraph_only(para_docx, style)
            continue

        _apply_full(para_docx, style, force_black=force_black)


def _apply_word_style(para, role: Role, docx_doc: DocxDocument) -> bool:
    """Attach the document's built-in style for ``role``; True if applied.

    A document that does not define the style is left alone. Inventing a
    reference to a style the file does not contain is worse than doing
    nothing: Word would render it unpredictably.
    """
    style_id = _find_word_style_id(docx_doc, role)
    if style_id is None:
        return False

    pPr = para._p.get_or_add_pPr()
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        pStyle = OxmlElement("w:pStyle")
        pPr.insert(0, pStyle)  # pStyle comes first in the pPr sequence
    pStyle.set(qn("w:val"), style_id)
    return True


def _find_word_style_id(docx_doc: DocxDocument, role: Role) -> str | None:
    aliases = _WORD_STYLE_ALIASES.get(role)
    if aliases is None:
        return None
    wanted = {_normalise_style_name(alias) for alias in aliases}

    for style in docx_doc.styles.element.findall(qn("w:style")):
        if style.get(qn("w:type")) != "paragraph":
            continue
        name_el = style.find(qn("w:name"))
        if name_el is None:
            continue
        name = name_el.get(qn("w:val")) or ""
        if _normalise_style_name(name) in wanted:
            return style.get(qn("w:styleId"))
    return None


def _normalise_style_name(name: str) -> str:
    return name.strip().lower().replace(" ", "")


def _apply_full(para, style: ParagraphStyle, *, force_black: bool = False) -> None:
    """Apply paragraph + run properties from ``style``."""
    _apply_paragraph_only(para, style)

    for run in para.runs:
        _apply_run(run, style, force_black=force_black)


def _apply_paragraph_only(para, style: ParagraphStyle) -> None:
    """Apply only paragraph-level properties (alignment, spacing, indent)."""
    pf = para.paragraph_format

    if style.alignment is not None:
        para.alignment = _ALIGN_MAP[style.alignment]

    if style.line_spacing is not None:
        pf.line_spacing = style.line_spacing

    if style.space_before_pt is not None:
        pf.space_before = Pt(style.space_before_pt)

    if style.space_after_pt is not None:
        pf.space_after = Pt(style.space_after_pt)

    if style.first_line_indent_chars is not None:
        char_pt = style.font_size_pt if style.font_size_pt is not None else 12.0
        pf.first_line_indent = Pt(char_pt * style.first_line_indent_chars)

    if style.left_indent_pt is not None:
        pf.left_indent = Pt(style.left_indent_pt)


def _apply_run(run, style: ParagraphStyle, *, force_black: bool = False) -> None:
    """Apply font name, size, weight, italic, color to a single run."""
    if style.font_size_pt is not None:
        run.font.size = Pt(style.font_size_pt)
    if style.bold is not None:
        run.font.bold = style.bold
    if style.italic is not None:
        run.font.italic = style.italic
    if style.color_hex is not None:
        run.font.color.rgb = RGBColor.from_string(style.color_hex)
    elif force_black:
        run.font.color.rgb = RGBColor(0, 0, 0)

    if style.font_name is not None:
        run.font.name = style.font_name

    if style.font_name_east_asian is not None:
        _set_east_asian_font(run, style.font_name_east_asian)


def _set_east_asian_font(run, font_name: str) -> None:
    """Set ``w:rFonts/@w:eastAsia`` on the run's ``rPr``.

    python-docx's ``run.font.name`` only writes ``w:ascii`` and
    ``w:hAnsi``; East Asian text reads ``w:eastAsia``.
    """
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(f"{{{W_NS}}}rFonts")
    if rFonts is None:
        rFonts = etree.SubElement(rPr, f"{{{W_NS}}}rFonts")
    rFonts.set(f"{{{W_NS}}}eastAsia", font_name)
    if rFonts.get(f"{{{W_NS}}}ascii") is None:
        rFonts.set(f"{{{W_NS}}}ascii", font_name)
    if rFonts.get(f"{{{W_NS}}}hAnsi") is None:
        rFonts.set(f"{{{W_NS}}}hAnsi", font_name)
